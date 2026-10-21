"""Bulk MACD from an Excel (.xlsx) or CSV sheet - one change per row."""

import csv
import os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from axl_client import AXLError

ACTIONS = {
    "ADD_PHONE": "New phone/Jabber device (+ line created if missing, user associated)",
    "UPDATE_PHONE": "Change description / device pool / CSS / location / template / owner",
    "REPLACE_PHONE": "Hardware swap: rename DeviceName to NewDeviceName (new MAC)",
    "MOVE_PHONE": "Site move: new DevicePool / Location / CSS",
    "ADD_LINE": "Add DN to an existing device (creates the DN if missing)",
    "MOVE_LINE": "Move DN from DeviceName to TargetDevice",
    "SHARE_LINE": "Add DN to TargetDevice and keep it on DeviceName",
    "UPDATE_LINE": "Change alerting name / description / CFA / line CSS / VM / NewDN",
    "REMOVE_LINE": "Remove DN from DeviceName (DN itself kept)",
    "ASSOCIATE_USER": "Associate DeviceName to OwnerUserID (+ primary extension = DN)",
    "DELETE_PHONE": "Delete device (DeleteOrphanLines=Y also deletes unused DNs)",
    "DELETE_LINE": "Delete DN (only if on no device)",
    "RESTART_PHONE": "Restart device",
}

ACTIONS_AR = {
    "ADD_PHONE": "هاتف / جهاز Jabber جديد (يُنشأ الرقم إن لم يكن موجوداً ويُربط بالمستخدم)",
    "UPDATE_PHONE": "تعديل الوصف / مجموعة الأجهزة / CSS / الموقع / القالب / المالك",
    "REPLACE_PHONE": "استبدال الجهاز: تغيير DeviceName إلى NewDeviceName (MAC جديد)",
    "MOVE_PHONE": "نقل موقع: DevicePool / Location / CSS جديدة",
    "ADD_LINE": "إضافة رقم إلى جهاز موجود (يُنشأ الرقم إن لم يكن موجوداً)",
    "MOVE_LINE": "نقل الرقم من DeviceName إلى TargetDevice",
    "SHARE_LINE": "إضافة الرقم إلى TargetDevice مع إبقائه على DeviceName",
    "UPDATE_LINE": "تعديل اسم التنبيه / الوصف / التحويل / CSS الخط / البريد الصوتي / NewDN",
    "REMOVE_LINE": "إزالة الرقم من DeviceName (مع الاحتفاظ بالرقم)",
    "ASSOCIATE_USER": "ربط DeviceName بـ OwnerUserID (والرقم الأساسي = DN)",
    "DELETE_PHONE": "حذف الجهاز (DeleteOrphanLines=Y يحذف أيضاً الأرقام غير المستخدمة)",
    "DELETE_LINE": "حذف الرقم (فقط إذا لم يكن على أي جهاز)",
    "RESTART_PHONE": "إعادة تشغيل الجهاز",
}

NOTES_AR = [
    "الصفوف في ورقة MACD أمثلة فقط - استبدلها بتغييراتك.",
    "الخلايا الفارغة يتم تجاهلها (لا يتم مسحها). يجب أن تطابق الأسماء ما في CUCM تماماً.",
    "القسم (Partition) فارغ = القسم <None>.",
    "استخدم 'تحقق' أولاً: يفحص وجود الأجهزة والأرقام والمستخدمين دون تغيير أي شيء.",
    "أسماء الأعمدة والإجراءات تبقى بالإنجليزية لأن الأداة تقرأها.",
]

COLUMNS = ["Action", "DeviceName", "Product", "Protocol", "DevicePool", "CSS", "Location",
           "ButtonTemplate", "SecurityProfile", "SIPProfile", "Description", "OwnerUserID",
           "DN", "Partition", "AlertingName", "LineDescription", "LineCSS", "VoicemailProfile",
           "CFA_Destination", "CFA_CSS", "E164Mask", "NewDeviceName", "TargetDevice",
           "NewDN", "NewPartition", "DeleteOrphanLines"]

EXAMPLES = [
    ["ADD_PHONE", "SEP001122334455", "Cisco 8845", "SIP", "DP_HQ", "CSS_Internal", "Hub_None",
     "Standard 8845 SIP", "Cisco 8845 - Standard SIP Non-Secure Profile", "Standard SIP Profile",
     "John Smith 8845", "jsmith", "4001", "PT_Internal", "John Smith", "John Smith",
     "CSS_Line_Intl", "", "", "", "+96824XXXXXX", "", "", "", "", ""],
    ["ADD_PHONE", "CSFJSMITH", "Cisco Unified Client Services Framework", "SIP", "DP_HQ",
     "CSS_Internal", "", "Standard Client Services Framework", "", "", "John Smith Jabber",
     "jsmith", "4001", "PT_Internal", "John Smith", "", "", "", "", "", "", "", "", "", "", ""],
    ["UPDATE_LINE", "", "", "", "", "", "", "", "", "", "", "", "4001", "PT_Internal",
     "John A. Smith", "", "", "", "5001", "CSS_Internal", "", "", "", "", "", ""],
    ["REPLACE_PHONE", "SEP001122334455", "", "", "", "", "", "", "", "", "", "", "", "", "", "",
     "", "", "", "", "", "SEP00AABBCCDDEE", "", "", "", ""],
    ["MOVE_LINE", "SEP00AABBCCDDEE", "", "", "", "", "", "", "", "", "", "", "4001", "PT_Internal",
     "", "", "", "", "", "", "", "", "SEP0099887766FF", "", "", ""],
    ["DELETE_PHONE", "SEP0099887766FF", "", "", "", "", "", "", "", "", "", "", "", "", "", "",
     "", "", "", "", "", "", "", "", "", "Y"],
]


def write_template(path, lang="en"):
    wb = Workbook()
    ws = wb.active
    ws.title = "MACD"
    ws.append(COLUMNS)
    for row in EXAMPLES:
        ws.append(row)
    head = PatternFill("solid", fgColor="1F4E78")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head
        c.alignment = Alignment(horizontal="center")
    for i, col in enumerate(COLUMNS, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(14, len(col) + 4)
    ws.freeze_panes = "B2"
    dv = DataValidation(type="list", formula1='"' + ",".join(ACTIONS) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("A2:A2000")

    help_ws = wb.create_sheet("Help")
    arabic = lang == "ar"
    help_ws.append(["Action", "What it does"] + (["الوصف"] if arabic else []))
    for k, v in ACTIONS.items():
        help_ws.append([k, v] + ([ACTIONS_AR[k]] if arabic else []))
    help_ws.append([])
    help_ws.append(["Notes"])
    help_ws.append(["Example rows in MACD are samples - replace them with your changes."])
    help_ws.append(["Blank cells are ignored (not cleared). Names must match CUCM exactly."])
    help_ws.append(["Partition blank = <None> partition."])
    help_ws.append(["Use 'Validate' first: it checks devices/DNs/users exist without changing anything."])
    if arabic:
        help_ws.append([])
        help_ws.append(["", "", "ملاحظات"])
        for n in NOTES_AR:
            help_ws.append(["", "", n])
        for row in help_ws.iter_rows(min_col=3, max_col=3):
            row[0].alignment = Alignment(horizontal="right", wrap_text=True, readingOrder=2)
        help_ws.column_dimensions["C"].width = 80
    for c in help_ws[1]:
        c.font = Font(bold=True)
    help_ws.column_dimensions["A"].width = 18
    help_ws.column_dimensions["B"].width = 90
    wb.save(path)


def read_rows(path):
    if path.lower().endswith(".csv"):
        with open(path, newline="", encoding="utf-8-sig") as f:
            rows = list(csv.DictReader(f))
    else:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb["MACD"] if "MACD" in wb.sheetnames else wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(it)]
        rows = [dict(zip(header, r)) for r in it]
    out = []
    for n, r in enumerate(rows, start=2):
        clean = {k: ("" if v is None else str(v).strip()) for k, v in r.items() if k}
        # Excel turns DNs like 4001 into 4001.0
        for k in ("DN", "NewDN", "CFA_Destination"):
            if clean.get(k, "").endswith(".0") and clean[k][:-2].isdigit():
                clean[k] = clean[k][:-2]
        if any(clean.values()):
            clean["_row"] = n
            clean["Action"] = clean.get("Action", "").upper()
            out.append(clean)
    return out


def _req(row, *cols):
    missing = [c for c in cols if not row.get(c)]
    if missing:
        raise AXLError(f"missing {', '.join(missing)}")


def validate_row(macd, r):
    """Read-only checks. Returns a short note; raises AXLError on a problem."""
    a = r["Action"]
    if a not in ACTIONS:
        raise AXLError(f"unknown action '{a}'")
    notes = []
    if a == "ADD_PHONE":
        _req(r, "DeviceName", "Product", "DevicePool")
        if macd.phone_exists(r["DeviceName"]):
            raise AXLError(f"{r['DeviceName']} already exists")
        if r.get("DN"):
            notes.append("DN exists - will reuse" if macd.line_exists(r["DN"], r.get("Partition"))
                         else "DN will be created")
    elif a in ("UPDATE_PHONE", "REPLACE_PHONE", "MOVE_PHONE", "REMOVE_LINE", "DELETE_PHONE",
               "RESTART_PHONE", "ASSOCIATE_USER", "ADD_LINE"):
        _req(r, "DeviceName")
        if not macd.phone_exists(r["DeviceName"]):
            raise AXLError(f"device {r['DeviceName']} not found")
        if a == "REPLACE_PHONE":
            _req(r, "NewDeviceName")
            if macd.phone_exists(r["NewDeviceName"]):
                raise AXLError(f"{r['NewDeviceName']} already exists")
        if a in ("REMOVE_LINE", "ADD_LINE"):
            _req(r, "DN")
        if a == "ASSOCIATE_USER":
            _req(r, "OwnerUserID")
    elif a in ("MOVE_LINE", "SHARE_LINE"):
        _req(r, "DN", "TargetDevice")
        if not macd.phone_exists(r["TargetDevice"]):
            raise AXLError(f"target {r['TargetDevice']} not found")
        if not macd.line_exists(r["DN"], r.get("Partition")):
            raise AXLError(f"DN {r['DN']}/{r.get('Partition') or '<None>'} not found")
    elif a in ("UPDATE_LINE", "DELETE_LINE"):
        _req(r, "DN")
        if not macd.line_exists(r["DN"], r.get("Partition")):
            raise AXLError(f"DN {r['DN']}/{r.get('Partition') or '<None>'} not found")
        if a == "DELETE_LINE":
            n = macd.line_device_count(r["DN"], r.get("Partition"))
            if n:
                raise AXLError(f"DN still on {n} device(s)")
    if r.get("OwnerUserID") and a in ("ADD_PHONE", "UPDATE_PHONE", "ASSOCIATE_USER"):
        try:
            macd.get_user(r["OwnerUserID"])
        except AXLError:
            raise AXLError(f"user {r['OwnerUserID']} not found")
    return "; ".join(notes) or "OK"


def _opt(r, col):
    v = r.get(col, "")
    return v if v else None


def run_row(macd, r):
    a = r["Action"]
    dev, dn, pt = r.get("DeviceName", ""), r.get("DN", ""), r.get("Partition", "")
    if a == "ADD_PHONE":
        name = macd.add_phone({
            "name": dev, "product": r["Product"], "protocol": r.get("Protocol") or "SIP",
            "devicePool": r["DevicePool"], "css": r.get("CSS"), "location": r.get("Location"),
            "buttonTemplate": r.get("ButtonTemplate"), "securityProfile": r.get("SecurityProfile"),
            "sipProfile": r.get("SIPProfile"), "description": r.get("Description"),
            "owner": r.get("OwnerUserID"), "dn": dn, "partition": pt,
            "alerting": r.get("AlertingName", ""), "lineDescription": r.get("LineDescription", ""),
            "lineCss": r.get("LineCSS", ""), "vmProfile": r.get("VoicemailProfile", ""),
            "cfaDest": r.get("CFA_Destination", ""), "cfaCss": r.get("CFA_CSS", ""),
            "e164Mask": r.get("E164Mask", "")})
        return f"created {name}"
    if a == "UPDATE_PHONE":
        macd.update_phone(dev, {"description": _opt(r, "Description"),
                                "devicePoolName": _opt(r, "DevicePool"),
                                "callingSearchSpaceName": _opt(r, "CSS"),
                                "locationName": _opt(r, "Location"),
                                "phoneTemplateName": _opt(r, "ButtonTemplate"),
                                "securityProfileName": _opt(r, "SecurityProfile"),
                                "sipProfileName": _opt(r, "SIPProfile"),
                                "ownerUserName": _opt(r, "OwnerUserID")})
        return "updated"
    if a == "REPLACE_PHONE":
        return "renamed to " + macd.update_phone(dev, {"newName": r["NewDeviceName"]})
    if a == "MOVE_PHONE":
        macd.move_phone(dev, _opt(r, "DevicePool"), _opt(r, "Location"), _opt(r, "CSS"),
                        _opt(r, "Description"))
        return "moved"
    if a == "ADD_LINE":
        macd.ensure_line(dn, pt, r.get("AlertingName", ""), r.get("LineDescription", ""),
                         r.get("LineCSS", ""), r.get("VoicemailProfile", ""))
        idx = macd.add_line_to_phone(dev, dn, pt, {"alerting": r.get("AlertingName", ""),
                                                   "e164Mask": r.get("E164Mask", "")})
        macd.apply_phone(dev)
        return f"added as line {idx}"
    if a in ("MOVE_LINE", "SHARE_LINE"):
        macd.move_line(dn, pt, dev, r["TargetDevice"], keep_shared=(a == "SHARE_LINE"))
        return "done"
    if a == "UPDATE_LINE":
        macd.update_line(dn, pt, {
            "description": _opt(r, "LineDescription"), "alertingName": _opt(r, "AlertingName"),
            "shareLineAppearanceCssName": _opt(r, "LineCSS"),
            "voiceMailProfileName": _opt(r, "VoicemailProfile"),
            "cfaDest": _opt(r, "CFA_Destination"), "cfaCss": _opt(r, "CFA_CSS"),
            "newPattern": _opt(r, "NewDN"), "newPartition": _opt(r, "NewPartition")})
        return "updated"
    if a == "REMOVE_LINE":
        macd.remove_line_from_phone(dev, dn, pt)
        macd.apply_phone(dev)
        return "removed"
    if a == "ASSOCIATE_USER":
        macd.associate_device_to_user(r["OwnerUserID"], dev, dn, pt)
        return "associated"
    if a == "DELETE_PHONE":
        gone = macd.delete_phone(dev, delete_orphan_lines=r.get("DeleteOrphanLines", "").upper()
                                 in ("Y", "YES", "TRUE", "1"))
        return "deleted" + (f" (+ DNs {', '.join(gone)})" if gone else "")
    if a == "DELETE_LINE":
        macd.delete_line(dn, pt)
        return "deleted"
    if a == "RESTART_PHONE":
        macd.restart_phone(dev)
        return "restarted"
    raise AXLError(f"unknown action '{a}'")


def write_results(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    cols = ["Row", "Action", "DeviceName", "DN", "Status", "Message"]
    ws.append(cols)
    for r in rows:
        ws.append([r.get("_row"), r.get("Action"), r.get("DeviceName"), r.get("DN"),
                   r.get("_status"), r.get("_message")])
    for c in ws[1]:
        c.font = Font(bold=True)
    red = PatternFill("solid", fgColor="F8CBAD")
    green = PatternFill("solid", fgColor="C6EFCE")
    for row in ws.iter_rows(min_row=2):
        row[4].fill = green if row[4].value in ("OK", "DONE") else red
    for col, w in zip("ABCDEF", (6, 16, 22, 12, 10, 80)):
        ws.column_dimensions[col].width = w
    wb.save(path)
    return os.path.abspath(path)
